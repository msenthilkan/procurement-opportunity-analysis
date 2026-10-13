"""Generate synthetic sample input files with the same layout as the real exports.

All suppliers, people, emails and amounts are invented. Run:
    python scripts/generate_sample_data.py
Files are written to sample-data/.
"""
import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

random.seed(42)
rng = random.Random(7)  # separate stream for P-card simulation; keeps the original sample unchanged
OUT = Path(__file__).resolve().parent.parent / "sample-data"
OUT.mkdir(exist_ok=True)

# ---------------------------------------------------------------- taxonomy
# (L1 family, L2 category, L3 long, L3 short, led by, mean line value, lines weight, catalogue share, P-card policy)
# P-card policy: Allowed = card is the preferred route for low-value tail buys;
# Restricted = card only within limits/approved merchants (catalogue or IT route first);
# Blocked = never on card (services with tax, IP, co-employment or contract risk, or a dedicated channel).
TAXONOMY = [
    ("Sales, Communication & Marketing", "Media", "Media Buying - Digital and TV", "Media buying", "Business Led", 45000, 6, 0.0, "Restricted"),
    ("Sales, Communication & Marketing", "Media", "Market Research and Consumer Insights", "Market research", "Business Led", 18000, 3, 0.0, "Allowed"),
    ("Sales, Communication & Marketing", "Activation", "Events and Sponsorship", "Events", "Business Led", 22000, 3, 0.0, "Restricted"),
    ("Sales, Communication & Marketing", "Activation", "Promotional and POS Materials", "POS materials", "Procurement Led", 2500, 10, 0.25, "Allowed"),
    ("Maintenance, Repairs & Operations", "MRO Materials", "Spare Parts - Production Equipment", "Spare parts", "Business Led", 1800, 12, 0.25, "Allowed"),
    ("Maintenance, Repairs & Operations", "MRO Materials", "Industrial Consumables", "Consumables", "Procurement Led", 350, 14, 0.45, "Allowed"),
    ("Maintenance, Repairs & Operations", "MRO Services", "Equipment Maintenance Services", "Maintenance services", "Business Led", 6500, 5, 0.0, "Restricted"),
    ("Facilities Management", "Soft Services", "Cleaning Services", "Cleaning", "Procurement Led", 4200, 5, 0.0, "Blocked"),
    ("Facilities Management", "Soft Services", "Security Services", "Security", "Procurement Led", 5200, 3, 0.0, "Blocked"),
    ("Facilities Management", "Soft Services", "Uniforms and Laundry", "Uniforms", "Procurement Led", 600, 8, 0.40, "Restricted"),
    ("Transport, Logistics, Warehousing", "Logistics", "Freight Forwarding", "Freight", "Procurement Led", 9500, 6, 0.0, "Allowed"),
    ("Transport, Logistics, Warehousing", "Logistics", "Third-party Warehousing", "Warehousing", "Procurement Led", 12000, 3, 0.0, "Blocked"),
    ("Information & Communications Technology", "IT", "Software Licences and SaaS", "Software", "Business Led", 8000, 4, 0.0, "Restricted"),
    ("Information & Communications Technology", "IT", "IT Hardware and Peripherals", "IT hardware", "Procurement Led", 900, 9, 0.50, "Restricted"),
    ("Information & Communications Technology", "Telecom", "Mobile and Fixed Telecom", "Telecom", "Business Led", 3000, 3, 0.0, "Blocked"),
    ("Professional Services", "Advisory", "Management Consulting", "Consulting", "Business Led", 25000, 2, 0.0, "Blocked"),
    ("Professional Services", "Advisory", "Legal Services", "Legal", "Business Led", 9000, 2, 0.0, "Blocked"),
    ("Travel, Meetings & Events", "Travel", "Hotels and Accommodation", "Hotels", "Procurement Led", 700, 9, 0.0, "Blocked"),
    ("Travel, Meetings & Events", "Meetings", "Meeting Venues", "Venues", "Procurement Led", 3500, 3, 0.0, "Allowed"),
    ("HR Services", "Workforce", "Contingent Labour", "Contingent labour", "Procurement Led", 7000, 4, 0.0, "Blocked"),
    ("HR Services", "Workforce", "Training and Development", "Training", "Business Led", 2500, 3, 0.0, "Allowed"),
    ("Office", "Office Supplies", "Office Supplies and Stationery", "Office supplies", "Procurement Led", 150, 16, 0.65, "Restricted"),
]
UNMAPPED = ["Miscellaneous Services", "Other Goods"]  # appear in spend, not in taxonomy

COUNTRIES = {"US01": 30, "AE01": 14, "DE01": 10, "TR01": 10, "ZA01": 7, "CA01": 6, "ID01": 5, "JP01": 2}

FIRST = ["Alex", "Sam", "Jordan", "Taylor", "Morgan", "Casey", "Riley", "Jamie", "Avery", "Quinn",
         "Drew", "Rowan", "Parker", "Reese", "Hayden", "Emerson", "Skyler", "Finley", "Kai", "Logan"]
LAST = ["Smith", "Khan", "Garcia", "Okafor", "Novak", "Tanaka", "Silva", "Meyer", "Yilmaz", "Patel",
        "Dubois", "Rossi", "Nguyen", "Cohen", "Larsen"]


def bold_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True)


def save(name, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = name
    ws.append(headers)
    for r in rows:
        ws.append(r)
    bold_header(ws)
    wb.save(OUT / f"{name}.xlsx")


def us_date(d):
    return f"{d.month}/{d.day}/{d.year}"


# ---------------------------------------------------------------- taxonomy file
tax_rows = []
for i, (l1, l2, l3, short, led, _mean, _w, _cat, pcard) in enumerate(TAXONOMY, start=1):
    tax_rows.append(["Indirect", f"F{i:02d}", l1, f"C{i:03d}", l2, f"CM{i:04d}", short, l3,
                     f"Synthetic description for {short.lower()}", led,
                     "Centre-led" if led.startswith("Procurement") else "Local",
                     "Strategic" if i % 3 == 0 else "Tactical",
                     "Contract" if i % 2 else "Catalogue", pcard])
# Scenario commodities (scripted below; see docs/scenarios.md). Two rows share the short name
# "Lab Testing" on purpose (ambiguous match); "Laboratory Supplies" is not in the scope matrix.
# (L1 family, L2 category, L3 long, L3 short, led by, P-card policy)
SCENARIO_TAX = [
    ("Facilities Management", "Workwear", "Protective Workwear", "Workwear", "Procurement Led", "Restricted"),
    ("Transport, Logistics, Warehousing", "Courier & Postal", "Courier Services", "Courier", "Procurement Led", "Restricted"),
    ("Facilities Management", "Hygiene", "Pest Control Services", "Pest control", "Procurement Led", "Blocked"),
    ("Sales, Communication & Marketing", "Digital Marketing", "Digital Analytics Services", "Digital analytics", "Business Led", "Blocked"),
    ("Professional Services", "Language Services", "Translation Services", "Translation", "Procurement Led", "Restricted"),
    ("Professional Services", "Strategy Advisory", "Corporate Strategy Advisory", "Strategy advisory", "Business Led", "Blocked"),
    ("Sales, Communication & Marketing", "Trade Investment", "Customer Rebates", "Rebates", "Business Led", "Blocked"),
    ("Real Estate", "Leases", "Office Rent", "Rent", "Business Led", "Blocked"),
    ("Research & Quality", "Lab Supplies", "Laboratory Supplies", "Lab supplies", "Procurement Led", "Allowed"),
    ("Research & Quality", "Testing", "Lab Testing - Microbiology", "Lab Testing", "Business Led", "Restricted"),
    ("Research & Quality", "Testing", "Lab Testing - Chemistry", "Lab Testing", "Business Led", "Restricted"),
]
for j, (l1, l2, l3, short, led, pcard) in enumerate(SCENARIO_TAX, start=len(TAXONOMY) + 1):
    tax_rows.append(["Indirect", f"F{j:02d}", l1, f"C{j:03d}", l2, f"CM{j:04d}", short, l3,
                     f"Synthetic description for {short.lower()}", led,
                     "Centre-led" if led.startswith("Procurement") else "Local", "Tactical", "Contract", pcard])
save("qTaxonomy",
     ["Level 0 (Group)", "Purchasing Family number", "Level 1 (Family)", "Purchasing Category number",
      "Level 2 (Category)", "Purchasing Commodity number", "Level 3 (Commodity) - Short",
      "Level 3 (Commodity) - Long", "Description", "Procurement Or Business Led",
      "Operating Model Types", "Procurement Category Classification", "Execution Strategy",
      "P-card Policy"],
     tax_rows)

# ---------------------------------------------------------------- people
people = []
used = set()
while len(people) < 60:
    n = f"{random.choice(FIRST)} {random.choice(LAST)}"
    if n in used:
        continue
    used.add(n)
    people.append(n)
procurement = people[:10]
email = {p: p.lower().replace(" ", ".") + "@example.com" for p in people}

um_rows = []
for i, p in enumerate(people, start=1):
    um_rows.append(["No", f"U{i:04d}", p, "Employee", "Yes", "Yes", "2026-09-30", "", "Chrome", email[p]])
# an inactive duplicate account to show the "active first" rule
um_rows.append(["Yes", "U9999", people[12], "Employee", "No", "No", "2023-01-15", "", "Edge", "old." + email[people[12]]])
save("qUserMaster",
     ["Locked", "User ID", "Name", "Type", "Active", "Has Password", "Last Login", "Delegatee",
      "Last Browser Used", "Business Email Address"], um_rows)
save("tblProcurementGroup", ["Email"], [[email[p]] for p in procurement])

# ---------------------------------------------------------------- spend + approvals
spend_rows, appr_rows = [], []
weights = [t[6] for t in TAXONOMY]
start = dt.date(2025, 10, 1)
req_no = 1000
PCARD_LIMIT = 2500          # single-transaction limit used by the simulation
card_on_po = []             # Ariba POs paid by card: these also appear on the card statement
contracted_suppliers = {}   # L3 -> suppliers used on catalogue/contract lines
suppliers = {t[2]: [f"Supplier {random.randint(100, 999)}" for _ in range(random.randint(2, 12))] for t in TAXONOMY}

for _ in range(1400):
    req_no += 1
    rid = f"PR{req_no}"
    country = random.choices(list(COUNTRIES), weights=list(COUNTRIES.values()))[0]
    d = start + dt.timedelta(days=random.randint(0, 364))
    po_d = d + dt.timedelta(days=random.randint(0, 6))
    unmapped = random.random() < 0.03
    t = random.choices(TAXONOMY, weights=weights)[0]
    l3, short, led, mean, cat_share, pcard_policy = t[2], t[3], t[4], t[5], t[7], t[8]
    commodity = random.choice(UNMAPPED) if unmapped else random.choice([l3, short])  # short or long name
    supplier = random.choice(suppliers[l3])
    n_lines = random.randint(1, 4)

    led_proc = led.startswith("Procurement")
    approvers = random.sample(people[10:], random.randint(1, 2))
    if (led_proc and random.random() < 0.55) or random.random() < 0.15:
        approvers.append(random.choice(procurement))
    if random.random() < 0.02:
        approvers.append("Former Employee")  # not in user master

    line_types = []
    for ln in range(n_lines):
        catalog = random.random() < cat_share
        contract = "" if catalog else (f"CW{random.randint(1000, 9999)}" if random.random() < 0.12 else "")
        amt = round(max(20, random.lognormvariate(0, 0.8) * mean), 2)
        line_type = "Catalog Item" if catalog else ("Non-Catalog Item" if random.random() > 0.02 else "Service Item")
        line_types.append(line_type)
        # Some low-value non-catalogue POs are paid by P-card (Ariba "P-card on PO")
        pay = "Invoice"
        if not catalog and amt <= PCARD_LIMIT and pcard_policy != "Blocked" and rng.random() < 0.30:
            pay = "P-Card"
            card_on_po.append((f"PO{req_no}", country, d, supplier, amt, commodity))
        if catalog:
            contracted_suppliers.setdefault(l3, set()).add(supplier)
        elif contract:
            contracted_suppliers.setdefault(l3, set()).add(supplier)
        spend_rows.append([rid, supplier, us_date(d), country, contract, line_type, "Ordered",
                           f"GL {short.title()}", line_type, f"PO{req_no}", f"UNSPSC-{random.randint(10, 99)}",
                           commodity, f"ERP-{random.randint(100, 999)}", us_date(po_d),
                           round(random.uniform(0.5, 72), 1), 1, amt, amt, len(approvers), pay])
        for a in approvers:  # approver export: each line repeated per approver
            appr_rows.append([rid, us_date(d), country, commodity, line_type, contract, a, amt])

# Low-value tail: one-off small requisitions with many different suppliers (classic P-card candidates).
# Appended after the original sample so the original rows stay identical.
TAIL = [("Freight Forwarding", 160, 80), ("Training and Development", 90, 260), ("Meeting Venues", 60, 300),
        ("Market Research and Consumer Insights", 40, 320), ("Promotional and POS Materials", 60, 180)]
tax_by_l3 = {t[2]: t for t in TAXONOMY}
for l3, n, mean in TAIL:
    t = tax_by_l3[l3]
    for _ in range(n):
        req_no += 1
        rid = f"PR{req_no}"
        country = rng.choices(list(COUNTRIES), weights=list(COUNTRIES.values()))[0]
        d = start + dt.timedelta(days=rng.randint(0, 364))
        amt = round(min(480, max(15, rng.lognormvariate(0, 0.5) * mean)), 2)
        supplier = f"Supplier {rng.randint(1000, 9999)}"          # one-off supplier
        approver = rng.choice(people[10:])
        spend_rows.append([rid, supplier, us_date(d), country, "", "Non-Catalog Item", "Ordered",
                           f"GL {t[3].title()}", "Non-Catalog Item", f"PO{req_no}", f"UNSPSC-{rng.randint(10, 99)}",
                           rng.choice([t[2], t[3]]), f"ERP-{rng.randint(100, 999)}", us_date(d),
                           round(rng.uniform(0.5, 48), 1), 1, amt, amt, 1, "Invoice"])
        appr_rows.append([rid, us_date(d), country, t[2], "Non-Catalog Item", "", approver, amt])

# Scenario requisitions: each block makes one rule fire on purpose (docs/scenarios.md).
srng = random.Random(11)  # own stream; earlier rows are unchanged
scn_po_date = dt.date(2026, 3, 1)


def add_req(country, supplier, commodity, amount, catalog=False, contract="", proc_approver=False,
            approval="business"):
    """approval: "business" (+ procurement if proc_approver), "procurement" only, or "none"."""
    global req_no
    req_no += 1
    rid = f"PR{req_no}"
    d = scn_po_date + dt.timedelta(days=srng.randint(0, 180))
    line_type = "Catalog Item" if catalog else "Non-Catalog Item"
    if approval == "procurement":
        approvers = [procurement[1]]
    elif approval == "none":
        approvers = []
    else:
        approvers = [srng.choice(people[10:])] + ([procurement[0]] if proc_approver else [])
    amt = round(amount * srng.uniform(0.85, 1.15), 2)
    spend_rows.append([rid, supplier, us_date(d), country, contract, line_type, "Ordered",
                       f"GL {commodity}", line_type, f"PO{req_no}", f"UNSPSC-{srng.randint(10, 99)}",
                       commodity, f"ERP-{srng.randint(100, 999)}", us_date(d),
                       round(srng.uniform(0.5, 48), 1), 1, amt, amt, len(approvers), "Invoice"])
    for a in approvers:
        appr_rows.append([rid, us_date(d), country, commodity, line_type, contract, a, amt])


def add_many(n, *args, **kw):
    for _ in range(n):
        add_req(*args, **kw)


# S-LEAK: catalogue with WW-1 in UAE, yet UAE still buys from WW-1 off-catalogue
add_many(15, "AE01", "Supplier WW-1", "Protective Workwear", 120, catalog=True)
add_many(10, "AE01", "Supplier WW-1", "Protective Workwear", 150, proc_approver=True)
add_many(2, "AE01", "Supplier WW-2", "Protective Workwear", 140)
# S-EXTEND: CR-1 is on catalogue in USA only; Turkey buys from CR-1 off-catalogue
add_many(20, "US01", "Supplier CR-1", "Courier Services", 60, catalog=True)
add_many(12, "TR01", "Supplier CR-1", "Courier Services", 70)
add_many(3, "TR01", "Supplier CR-3", "Courier Services", 65)
# S-REPLICATE: contract with PC-1 in Germany; Turkey and South Africa use other suppliers
add_many(12, "DE01", "Supplier PC-1", "Pest Control Services", 400, contract="CW-SCN-01")
add_many(8, "TR01", "Supplier PC-2", "Pest Control Services", 380)
add_many(6, "ZA01", "Supplier PC-3", "Pest Control Services", 350)
# S-REGIONAL: no agreement anywhere; DA-1 is the top supplier in USA, UAE and Turkey
add_many(6, "US01", "Supplier DA-1", "Digital Analytics Services", 9000)
add_many(1, "US01", "Supplier DA-4", "Digital Analytics Services", 5000)
add_many(4, "AE01", "Supplier DA-1", "Digital Analytics Services", 8000)
add_many(1, "AE01", "Supplier DA-5", "Digital Analytics Services", 4000)
add_many(3, "TR01", "Supplier DA-1", "Digital Analytics Services", 7000)
# S-NEW + S-OVERRIDE: no agreement, different supplier per country; many small lines
# (the frequency rule would say Catalogue, the scope matrix says rate-card contract)
add_many(16, "DE01", "Supplier TR-1", "Translation Services", 250, proc_approver=True)
add_many(10, "DE01", "Supplier TR-2", "Translation Services", 230)
add_many(4, "CA01", "Supplier TR-3", "Translation Services", 300)
# S-BUSINESS: tier C, chosen by leadership
add_many(2, "US01", "Supplier CS-1", "Corporate Strategy Advisory", 120000)
# S-EXCL-E: non-addressable spend raised on requisitions
for c in ("US01", "AE01", "DE01", "TR01"):
    add_many(6, c, f"Customer {c[:2]}-01", "Customer Rebates", 80000)
add_many(12, "DE01", "Landlord DE-1", "Office Rent", 35000)
add_many(12, "AE01", "Landlord AE-1", "Office Rent", 30000)
# S-NOTCLASS: in the taxonomy but not in the scope matrix
add_many(8, "DE01", "Supplier LS-1", "Laboratory Supplies", 600)
# S-AMBIG: short name shared by two taxonomy rows
add_many(3, "US01", "Supplier LT-1", "Lab Testing", 2500)
# S-APPR-PROC: approved by procurement only; S-APPR-NONE: no rows in the approver export
add_many(2, "AE01", "Supplier WW-1", "Protective Workwear", 160, approval="procurement")
add_many(1, "DE01", "Landlord DE-1", "Office Rent", 35000, approval="none")
# S-COUNTRY: purchasing organisation not in CountryMap
add_many(3, "BR01", "Supplier BR-1", "Office Supplies and Stationery", 200)

save("qSpend",
     ["[REQ] Requisition ID", "[REQ]Supplier (Common Supplier)", "[REQ]Requisition Date (Date)",
      "[REQ]Purchasing Company (Purchase Organization Id)", "[REQ]Contract (Contract)", "[REQ] Line Type",
      "[REQ] Requisition Status", "[REQ]GL Account (GL Account Name)", "[PO-P&I] Line Type", "[REQ] PO Id",
      "[REQ]Commodity (UNSPSC (L3))", "[REQ]Commodity (Commodity (L3))", "[REQ]ERP Commodity (ERP Commodity)",
      "[PO-P&I]Ordered Date (Date)", "max(Approval Time)", "sum(Line Count)", "sum(Requisition Spend)",
      "sum(PO Spend)", "sum(Approver Count)", "[PO-P&I] Payment Method"], spend_rows)
save("qApproverRaw",
     ["[REQ] Requisition ID", "[REQ]Requisition Date (Date)", "[REQ]Purchasing Company (Purchase Organization Id)",
      "[REQ]Commodity (Commodity)", "[REQ] Line Type", "[REQ]Contract (Contract)",
      "[APS]Approved by User (User)", "sum(Requisition Spend)"], appr_rows)

# ---------------------------------------------------------------- P-card simulation
# Merchant Category Codes (standard ISO 18245 codes) mapped to taxonomy L3.
# MCCs with no L3 carry their own policy (e.g. cash, gambling, bars are always blocked).
MCC_MAP = [
    (5111, "Stationery, office and school supply stores", "Office Supplies and Stationery", "Restricted"),
    (5943, "Stationery stores, office and school supply", "Office Supplies and Stationery", "Restricted"),
    (5085, "Industrial supplies", "Industrial Consumables", "Allowed"),
    (5251, "Hardware stores", "Industrial Consumables", "Allowed"),
    (5013, "Motor vehicle supplies and new parts", "Spare Parts - Production Equipment", "Allowed"),
    (7699, "Miscellaneous repair shops", "Equipment Maintenance Services", "Restricted"),
    (2741, "Miscellaneous publishing and printing", "Promotional and POS Materials", "Allowed"),
    (7333, "Commercial photography, art and graphics", "Promotional and POS Materials", "Allowed"),
    (7311, "Advertising services", "Media Buying - Digital and TV", "Restricted"),
    (8999, "Professional services not elsewhere classified", "Market Research and Consumer Insights", "Allowed"),
    (5811, "Caterers", "Events and Sponsorship", "Restricted"),
    (5045, "Computers, peripherals and software", "IT Hardware and Peripherals", "Restricted"),
    (5732, "Electronics stores", "IT Hardware and Peripherals", "Restricted"),
    (5734, "Computer software stores", "Software Licences and SaaS", "Restricted"),
    (7372, "Computer programming and data processing", "Software Licences and SaaS", "Restricted"),
    (4814, "Telecommunication services", "Mobile and Fixed Telecom", "Blocked"),
    (4215, "Courier services", "Freight Forwarding", "Allowed"),
    (4225, "Public warehousing and storage", "Third-party Warehousing", "Blocked"),
    (7349, "Cleaning, maintenance and janitorial services", "Cleaning Services", "Blocked"),
    (7393, "Detective, protective and security services", "Security Services", "Blocked"),
    (5137, "Uniforms and commercial clothing", "Uniforms and Laundry", "Restricted"),
    (7392, "Management, consulting and public relations", "Management Consulting", "Blocked"),
    (8111, "Legal services and attorneys", "Legal Services", "Blocked"),
    (7361, "Employment agencies and temporary help", "Contingent Labour", "Blocked"),
    (8299, "Schools and educational services", "Training and Development", "Allowed"),
    (7011, "Hotels, motels and resorts", "Hotels and Accommodation", "Blocked"),
    (5812, "Eating places and restaurants", "", "Restricted"),
    (5813, "Bars and taverns", "", "Blocked"),
    (6011, "Automated cash disbursements", "", "Blocked"),
    (7995, "Betting and gambling", "", "Blocked"),
]
save("qMCCMap", ["MCC", "MCC Description", "Level 3 (Commodity) - Long", "MCC Policy"],
     [list(r) for r in MCC_MAP])

# (MCC, weight, mean amount): card spend is tail-heavy and low value
CARD_MIX = [(5111, 14, 90), (5943, 6, 70), (5085, 10, 180), (5251, 8, 120), (5013, 6, 260), (7699, 3, 420),
            (2741, 5, 300), (7333, 2, 450), (7311, 5, 650), (8999, 3, 500), (5811, 3, 380), (5045, 4, 700),
            (5732, 4, 450), (5734, 4, 220), (7372, 4, 300), (4814, 1, 160), (4215, 7, 85), (4225, 1, 900),
            (7349, 1, 350), (7393, 1, 600), (5137, 2, 140), (7392, 1, 1800), (8111, 1, 1500), (7361, 1, 1200),
            (8299, 6, 550), (7011, 4, 210), (5812, 5, 95), (5813, 1, 60), (6011, 1, 200), (7995, 1, 120)]
mcc_l3 = {m: l3 for m, _d, l3, _p in MCC_MAP}
short_of = {t[2]: t[3] for t in TAXONOMY}
merchants = {m: [f"Merchant {rng.randint(1000, 9999)}" for _ in range(rng.randint(4, 25))] for m, *_ in CARD_MIX}
cardholders = rng.sample(people[10:], 25)
holder_country = {h: rng.choices(list(COUNTRIES), weights=list(COUNTRIES.values()))[0] for h in cardholders}

card_rows = []
txn = 500000


def card_row(holder, country, d, merchant, mcc, amt, po_ref="", coded="", receipt=None):
    global txn
    txn += 1
    if receipt is None:
        receipt = "Yes" if rng.random() < 0.92 else "No"
    post = d + dt.timedelta(days=rng.randint(1, 3))
    card_rows.append([f"T{txn}", us_date(d), us_date(post), holder, email[holder], country, merchant, mcc,
                      round(amt, 2), receipt, po_ref, coded,
                      rng.choice(["Reconciled", "Reconciled", "Reconciled", "Pending"])])


# 1) Normal off-platform card transactions
for _ in range(1100):
    mcc, _w, mean = rng.choices(CARD_MIX, weights=[c[1] for c in CARD_MIX])[0]
    holder = rng.choice(cardholders)
    d = start + dt.timedelta(days=rng.randint(0, 364))
    l3 = mcc_l3[mcc]
    merchant = rng.choice(merchants[mcc])
    # ~8%: card used with a supplier that already has a catalogue or contract (leakage)
    if l3 and contracted_suppliers.get(l3) and rng.random() < 0.08:
        merchant = rng.choice(sorted(contracted_suppliers[l3]))
    amt = max(5, rng.lognormvariate(0, 0.7) * mean)
    if rng.random() < 0.03:  # ~3% above the single-transaction limit
        amt = PCARD_LIMIT * rng.uniform(1.05, 2.5)
    coded = short_of.get(l3, "") if (l3 and rng.random() < 0.4) else ""  # cardholder-coded commodity
    card_row(holder, holder_country[holder], d, merchant, mcc, amt, coded=coded)

# 2) Split transactions: same cardholder, merchant and day, each under the limit, together above it
for _ in range(20):
    mcc = rng.choice([5045, 5732, 7311, 5085, 2741])
    holder = rng.choice(cardholders)
    d = start + dt.timedelta(days=rng.randint(0, 364))
    merchant = rng.choice(merchants[mcc])
    for _k in range(rng.choice([2, 2, 3])):
        card_row(holder, holder_country[holder], d, merchant, mcc, PCARD_LIMIT * rng.uniform(0.6, 0.98))

# 3) Ariba POs paid by card: the same spend appears again on the statement, with the PO number
card_mcc_for = {}
for m, _d, l3, _p in MCC_MAP:
    if l3:
        card_mcc_for.setdefault(l3, m)
l3_of_name = {}
for t in TAXONOMY:
    l3_of_name[t[2]] = t[2]
    l3_of_name[t[3]] = t[2]
for po, country, d, supplier, amt, commodity in card_on_po:
    holder = rng.choice(cardholders)
    mcc = card_mcc_for.get(l3_of_name.get(commodity, ""), 5999)
    card_row(holder, country, d + dt.timedelta(days=rng.randint(0, 5)), supplier, mcc, amt,
             po_ref=po, receipt="Yes")

save("qPCard",
     ["Transaction ID", "Transaction Date", "Posting Date", "Cardholder Name", "Cardholder Email",
      "Purchase Org", "Merchant Name", "MCC", "Amount (Reporting Currency)", "Receipt Attached",
      "PO Reference", "Commodity (Cardholder Coded)", "Reconciliation Status"], card_rows)
print(f"Wrote {len(card_rows)} P-card transactions ({len(card_on_po)} are card-paid Ariba POs)")

# ---------------------------------------------------------------- scope matrix (Excel table)
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402
import sys  # noqa: E402
sys.path.insert(0, str(Path(__file__).resolve().parent))
from sample_scope_matrix import ROWS as SCOPE_ROWS  # noqa: E402

wb = Workbook()
ws = wb.active
ws.title = "tblScopeMatrix"
ws.append(["L3 Commodity", "Category", "Scope Tier", "Proposed Route", "Lead", "Practical solution"])
for r in SCOPE_ROWS:
    ws.append(list(r))
bold_header(ws)
tbl = Table(displayName="tblScopeMatrix", ref=f"A1:F{ws.max_row}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)
wb.save(OUT / "tblScopeMatrix.xlsx")

print(f"Wrote {len(spend_rows)} spend lines and {len(appr_rows)} approver rows to {OUT}")
